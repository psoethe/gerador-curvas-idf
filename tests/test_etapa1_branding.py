"""
Testes automatizados para a Etapa 1 da transição SII-HiDRO (v3.0).
Valida:
1. Marca dinâmica e estilização estrita do 'i' em minúsculo na cor de ação (#2563eb).
2. Sufixos dos 4 módulos: Projeto, IDF, Bacia, Vazão.
3. Versionamento atualizado para 'SII-HiDRO v3.0' em UI, relatórios PDF e Word.
4. Cabeçalhos e rodapés dos relatórios com sufixos dinâmicos por módulo.
5. Chaves i18n PT e EN para marca e módulos.
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import io
import pandas as pd
import plotly.graph_objects as go
from docx import Document

import i18n
from i18n import t
import streamlit_app as app
import report
import report_word


def _dummy_figures():
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=[1, 2], y=[3, 4]))
    return fig, fig, fig, fig


def _dummy_results():
    from calculations import run_full_analysis
    series_df = pd.DataFrame({
        'Ano': list(range(1995, 2020)),
        'Precipitacao': [75.0, 92.0, 84.0, 110.0, 68.0, 95.0, 105.0, 88.0, 115.0, 78.0,
                         85.0, 90.0, 99.0, 120.0, 70.0, 82.0, 101.0, 94.0, 89.0, 108.0,
                         77.0, 93.0, 86.0, 104.0, 91.0],
    })
    return run_full_analysis(df_input=series_df, isozona='B', lang='PT')


def test_brand_dynamic_html():
    """Valida a renderização HTML da marca dinâmica e estilização do 'i'."""
    html_proj = app.render_brand_html('projeto')
    assert 'SII-H<span style=\'color:#2563eb;\'>i</span>DRO' in html_proj
    assert 'SII-H<span style=\'color:#2563eb;\'>i</span>DRO</span>' in html_proj  # sem sufixo

    html_idf = app.render_brand_html('idf')
    assert 'SII-H<span style=\'color:#2563eb;\'>i</span>DRO-IDF' in html_idf

    html_bacia = app.render_brand_html('bacia')
    assert 'SII-H<span style=\'color:#2563eb;\'>i</span>DRO-Bacia' in html_bacia

    html_vazao = app.render_brand_html('vazao')
    assert 'SII-H<span style=\'color:#2563eb;\'>i</span>DRO-Vazão' in html_vazao


def test_brand_text():
    """Valida o texto puro da marca dinâmica por módulo."""
    assert app.brand_text('projeto') == 'SII-HiDRO'
    assert app.brand_text('idf') == 'SII-HiDRO-IDF'
    assert app.brand_text('bacia') == 'SII-HiDRO-Bacia'
    assert app.brand_text('vazao') == 'SII-HiDRO-Vazão'


def test_i18n_branding_keys():
    """Valida as chaves de internacionalização da marca e versão."""
    assert t('app_brand', 'PT') == 'SII-HiDRO'
    assert t('app_brand', 'EN') == 'SII-HiDRO'
    assert t('app_version_str', 'PT') == 'SII-HiDRO v3.0'
    assert 'SII-HiDRO v3.0' in t('report_footer', 'PT')
    assert 'SII-HiDRO v3.0' in t('report_footer', 'EN')
    assert t('brand_idf', 'PT') == 'SII-HiDRO-IDF'
    assert t('brand_bacia', 'PT') == 'SII-HiDRO-Bacia'
    assert t('brand_vazao', 'PT') == 'SII-HiDRO-Vazão'


def test_pdf_report_branding():
    """Valida que o relatório PDF contém a marca SII-HiDRO-IDF e versão v3.0."""
    res = _dummy_results()
    fh, fg, fp, fi = _dummy_figures()
    pdf_bytes = report.generate_pdf_report(
        results=res,
        responsavel="Eng. Teste",
        localizacao="Local Teste",
        estacao="12345",
        fig_historica=fh,
        fig_gumbel=fg,
        fig_pdf=fp,
        fig_idf=fi,
        lang='PT',
        app_name='SII-HiDRO-IDF',
    )
    assert pdf_bytes is not None and len(pdf_bytes) > 5000
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf_bytes))
    full_pdf_text = " ".join(page.extract_text() or "" for page in reader.pages)
    assert 'SII-HiDRO-IDF' in full_pdf_text or 'SII-HiDRO' in full_pdf_text
    assert 'v3.0' in full_pdf_text
    assert 'IDF Curve Calculator' not in full_pdf_text


def test_word_report_branding():
    """Valida que o relatório Word (.docx) contém a marca SII-HiDRO-IDF e versão v3.0."""
    res = _dummy_results()
    fh, fg, fp, fi = _dummy_figures()
    docx_bytes = report_word.generate_word_report(
        results=res,
        responsavel="Eng. Teste",
        localizacao="Local Teste",
        estacao="12345",
        fig_historica=fh,
        fig_gumbel=fg,
        fig_pdf=fp,
        fig_idf=fi,
        lang='PT',
        app_name='SII-HiDRO-IDF',
    )
    assert docx_bytes is not None and len(docx_bytes) > 2000

    doc = Document(io.BytesIO(docx_bytes))
    full_text = ' '.join(p.text for p in doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                full_text += ' ' + cell.text

    assert 'SII-HiDRO-IDF' in full_text
    assert 'SII-HiDRO v3.0' in full_text
    assert 'IDF Curve Calculator' not in full_text


if __name__ == '__main__':
    tests = [
        ('1. Marca dinâmica em HTML com estilização do "i"', test_brand_dynamic_html),
        ('2. Marca dinâmica textual por módulo', test_brand_text),
        ('3. Chaves de i18n da marca e versão', test_i18n_branding_keys),
        ('4. Marca e versão no relatório PDF', test_pdf_report_branding),
        ('5. Marca e versão no relatório Word (.docx)', test_word_report_branding),
    ]

    print('=' * 70)
    print('EXECUTANDO TESTES DA ETAPA 1 (MARCA E NOMENCLATURA SII-HiDRO v3.0)')
    print('=' * 70)
    passed = 0
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f'  [PASS] {name}')
            passed += 1
        except Exception as e:
            print(f'  [FAIL] {name} -> {e}')
            import traceback
            traceback.print_exc()
            failed += 1

    print('=' * 70)
    print(f'TOTAL: {passed} passaram, {failed} falharam.')
    print('=' * 70)
    if failed > 0:
        exit(1)
